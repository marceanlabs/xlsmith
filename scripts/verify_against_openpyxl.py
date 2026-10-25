#!/usr/bin/env python3
"""Check what xlsmith writes against openpyxl, cell by cell.

openpyxl is a different implementation from calamine, so agreement between them
is evidence about the values rather than a shared assumption. `data_only=True`
asks openpyxl for the value the file stores, which is what calamine gives: a
formula is only visible through its last saved result, and a workbook written by
a program that never calculated has no result to give.

    python3 scripts/verify_against_openpyxl.py            # check the committed fixtures
    python3 scripts/verify_against_openpyxl.py --rebuild   # regenerate them first
    python3 scripts/verify_against_openpyxl.py --quick     # skip the 200,000 row workbook
"""
import csv as csvmod
import datetime as dt
import io
import shutil
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "tests/fixtures"
BIN = ROOT / "target/release/xlsmith"


def build_simple() -> Path:
    """Write the small workbook, and the same data as .ods and .xls if LibreOffice is here."""
    from openpyxl import Workbook

    path = FIX / "simple.xlsx"
    wb = Workbook()
    sales = wb.active
    sales.title = "Sales"
    sales.append(["region", "units", "price", "date", "in stock"])
    for r in [
        ("North", 12, 9.99, dt.date(2026, 3, 1), True),
        ("South", 0, 4.5, dt.date(2026, 3, 2), False),
        ("East", 1456, 0.01, dt.date(2026, 3, 3), True),
        ("West", 3, 1200.0, dt.date(2026, 3, 4), True),
        ("North", 8, 15.25, dt.date(2026, 3, 5), False),
    ]:
        sales.append(r)
    notes = wb.create_sheet("Notes")
    notes.append(["key", "text"])
    notes.append(["a", "café — naïve"])
    notes.append(["b", 'comma, inside "quotes"'])
    notes.append(["c", None])
    notes.append(["d", "line\nbreak"])
    total = wb.create_sheet("Total")
    total["A1"] = "sum"
    total["B1"] = "=SUM(Sales!B2:B6)"
    wb.save(path)

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "ods", "--outdir", str(FIX), str(path)],
            check=False, capture_output=True,
        )
        subprocess.run(
            [soffice, "--headless", "--convert-to", "xls", "--outdir", str(FIX), str(path)],
            check=False, capture_output=True,
        )
    return path


def build_large(rows: int = 200_000) -> Path:
    from openpyxl import Workbook

    path = FIX / "large.xlsx"
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.append(["id", "name", "value", "date"])
    for i in range(rows):
        ws.append([i, f"row {i}", i * 1.5, dt.date(2026, 1, 1)])
    wb.save(path)
    return path


def xlsmith(*args: str) -> str:
    out = subprocess.run([str(BIN), *args], capture_output=True, text=True)
    if out.returncode != 0:
        print(f"FAILED: xlsmith {' '.join(args)}\n{out.stderr}", file=sys.stderr)
        sys.exit(1)
    return out.stdout


def rendered(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dt.date, dt.datetime)):
        d = value.date() if isinstance(value, dt.datetime) else value
        return d.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def trimmed(row: list[str]) -> list[str]:
    """Trailing empty fields carry no data; whether a writer emits them is style.

    A formula has no value until something calculates it, so a sheet whose only
    second cell is `=SUM(...)` is one column wide as far as the file's own values
    go, however openpyxl counts its cells.
    """
    out = list(row)
    while out and out[-1] == "":
        out.pop()
    return out


def check_values(path: Path, sheet: str) -> None:
    from openpyxl import load_workbook

    got = [trimmed(r) for r in csvmod.reader(io.StringIO(xlsmith("csv", str(path), "--sheet", sheet)))]
    ws = load_workbook(path, data_only=True)[sheet]
    expected = [trimmed([rendered(v) for v in row]) for row in ws.iter_rows(values_only=True)]
    if got != expected:
        for i, (g, e) in enumerate(zip(got, expected)):
            if g != e:
                print(
                    f"MISMATCH {path.name}/{sheet} row {i + 1}:\n  xlsmith:  {g}\n  openpyxl: {e}",
                    file=sys.stderr,
                )
                sys.exit(1)
        print(
            f"MISMATCH {path.name}/{sheet}: {len(got)} rows from xlsmith, "
            f"{len(expected)} from openpyxl",
            file=sys.stderr,
        )
        sys.exit(1)
    print(f"  {path.name}/{sheet}: {len(expected)} rows agree with openpyxl")


def main() -> None:
    if "--rebuild" in sys.argv:
        FIX.mkdir(parents=True, exist_ok=True)
        build_simple()

    print("values checked against openpyxl:")
    for sheet in ("Sales", "Notes", "Total"):
        check_values(FIX / "simple.xlsx", sheet)

    print("sheets command:")
    print("\n".join("  " + line for line in xlsmith("sheets", str(FIX / "simple.xlsx")).splitlines()))

    print("json with header (first lines):")
    print(
        "\n".join(
            "  " + line
            for line in xlsmith("json", str(FIX / "simple.xlsx"), "--header").splitlines()[:5]
        )
    )

    print("every sheet at once:")
    out_dir = Path("/tmp/xlsmith-sheets")
    out_dir.mkdir(exist_ok=True)
    for old in out_dir.glob("*"):
        old.unlink()
    xlsmith("csv", str(FIX / "simple.xlsx"), "--all-sheets", str(out_dir))
    print("  " + ", ".join(sorted(p.name for p in out_dir.glob("*"))))

    if "--quick" in sys.argv:
        print("quick mode: skipping the 200,000 row workbook")
        return

    large = build_large()
    start = time.monotonic()
    out = subprocess.run([str(BIN), "csv", str(large), "-o", "/tmp/large.csv"], capture_output=True)
    elapsed = time.monotonic() - start
    size = Path("/tmp/large.csv").stat().st_size
    status = "ok" if out.returncode == 0 else f"FAILED {out.stderr.decode()[:200]}"
    print(f"large: {elapsed:.2f}s for 200,000 rows -> {size / 1e6:.1f} MB ({status})")


if __name__ == "__main__":
    main()
